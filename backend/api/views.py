from django.http import HttpResponse, JsonResponse
from django.conf import settings
from django.db import transaction, connection
from django.http import JsonResponse, StreamingHttpResponse
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie, csrf_exempt
from django.utils.decorators import method_decorator
from django.db.models import Q
from django.contrib.auth import authenticate, login, logout, update_session_auth_hash
from django.views.decorators.http import require_POST, require_GET
from django.contrib.auth.decorators import login_required
from tablib import Dataset
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.views import APIView, View
from rest_framework.parsers import JSONParser
from rest_framework.response import Response
from rest_framework import status, generics
from rest_framework.pagination import PageNumberPagination
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework import status
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from datetime import datetime
from .resources import PersonResource
from .consumers import broadcast_to_crud01, broadcast_stats_update, broadcast_ws
from .models import Person, Log, Profile
from .serializers import PersonSerializer, LogSerializer
from datetime import datetime, timedelta
from urllib.parse import quote
import json
import urllib.parse
import os, io
import logging

logger = logging.getLogger(__name__)

# ✅ แจก CSRF token (frontend ต้องเรียกก่อน)
@ensure_csrf_cookie
def get_csrf_token(request):
    return JsonResponse({"detail": "CSRF cookie set"})

# ✅ Login
@require_POST  # 👈 2. บังคับให้รับเฉพาะ POST ป้องกัน JSONDecodeError จาก GET
def login_view(request):
    try:
        data = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({'detail': 'Empty or invalid JSON body'}, status=400)

    username = data.get("username")
    password = data.get("password")
    # 3. ดึงค่า expires_in มา (ยังใช้ได้กับ Session)
    expires_in = int(data.get("expires_in", 24 * 60 * 60))  # 1 วัน

    user = authenticate(request, username=username, password=password)
    
    if user is not None:
        # 4. Login ระบบ session ของ Django (จะสร้าง sessionid cookie)
        login(request, user)
        
        # 5. ตั้งค่าหมดอายุ session ตามที่ frontend ส่งมา
        request.session.set_expiry(expires_in)

        # 6. ดึงข้อมูล profile เพื่อส่งกลับ (เหมือน profile_view)
        if user.groups.filter(name='Dev').exists():
            status = 'Dev'
        else:
            status = 'Staff'

        # 7. ส่งข้อมูล user กลับไป (auth.js จะใช้ setUser(res.data))
        return JsonResponse({
            "username": user.username,
            "first_name": user.first_name,
            "last_name": user.last_name,
            "nickname": user.profile.nickname if hasattr(user, 'profile') else '',
            "status": status
        })

    else:
        return JsonResponse({"detail": "Invalid credentials"}, status=400)
    
@login_required
@require_POST
def change_password_view(request):
    try:
        data = json.loads(request.body)
        new_password = data.get('new_password')

        if not new_password:
            return JsonResponse({'error': 'New password not provided'}, status=400)

        user = request.user
        user.set_password(new_password) # เข้ารหัสและตั้งรหัสผ่านใหม่
        user.save()

        # อัปเดต session ของผู้ใช้เพื่อไม่ให้หลุดออกจากระบบ
        update_session_auth_hash(request, user)

        return JsonResponse({'message': 'Password updated successfully'}, status=200)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

# ✅ Logout
@require_POST
def logout_view(request):
    logout(request)
    return JsonResponse({"detail": "Logged out"})

# ✅ ดึง user ปัจจุบัน
@login_required
def profile_view(request):
    user = request.user # <-- ย้าย user มาไว้ข้างบนเพื่อใช้ร่วมกัน

    if request.method == 'GET':
        if user.groups.filter(name='Dev').exists():
            status = 'Dev'
        else:
            status = 'Staff'

        return JsonResponse({
            "username": user.username,
            "email": user.email,
            "first_name": user.first_name,
            "last_name": user.last_name,
            "nickname": user.profile.nickname if hasattr(user, 'profile') else '',
            "status": status
        })

    elif request.method == 'PUT':
        try:
            data = json.loads(request.body)

            user.first_name = data.get('first_name', user.first_name)
            user.last_name = data.get('last_name', user.last_name)
            user.save()
            profile, created = Profile.objects.get_or_create(user=user)
            if created:
                print(f"Created a new profile for user: {user.username}")

            profile.nickname = data.get('nickname', profile.nickname)
            profile.save()
                
            if user.groups.filter(name='Dev').exists():
                status = 'Dev'
            else:
                status = 'Staff'

            return JsonResponse({
                "username": user.username,
                "email": user.email,
                "first_name": user.first_name,
                "last_name": user.last_name,
                "nickname": profile.nickname, # ใช้ profile.nickname ที่เราเพิ่งบันทึก
                "status": status
            }, status=200)

        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)

    return JsonResponse({'error': 'Method not allowed'}, status=405)

def file_iterator(buffer, chunk_size=8192):
    buffer.seek(0)
    while True:
        chunk = buffer.read(chunk_size)
        if not chunk:
            break
        yield chunk

def get_filtered_persons(request):
    # 1. ดึงค่า (จะเป็น '0', '1', '2' หรือ None ถ้าไม่ส่งมา)
    verified_status = request.GET.get('verified_status', None)
    persons = Person.objects.all()

    if verified_status == '0':
        # "ยังไม่รายงานตัว" (code: 0) - ตรรกะจาก person_stats
        persons = persons.filter(verified1=0, verified2=0, verified3=0)
    
    elif verified_status == '1':
        # "รายงานตัวแล้ว" (code: 1) - ตรรกะจาก person_stats
        persons = persons.filter(Q(verified1=1) | Q(verified2=1) | Q(verified3=1))
        
    elif verified_status == '2':
        # "อยู่ในห้องพิธี" (code: 2) - ตรรกะจาก person_stats
        persons = persons.filter(Q(verified1=2) | Q(verified2=2) | Q(verified3=2))

    # ถ้า verified_status เป็น None (ทั้งหมด)
    # ก็จะไม่ทำอะไร (คืนค่า persons.all())
    
    return persons

def get_custom_sort_key(person):
    """
    Key สำหรับเรียงลำดับ 'เลขที่บัณฑิต' (varchar)
    ลำดับ: ป > ปท > ตัวเลข
    """
    person_id = str(person.id) # 'id' คือ 'เลขที่บัณฑิต'

    if person_id.startswith('ปท'):
        group = 1  # 1. กลุ่ม "ปท"
        num_part = person_id[2:]
    elif person_id.startswith('ป'):
        group = 0  # 0. กลุ่ม "ป" (มาก่อน)
        num_part = person_id[1:]
    elif person_id.isdigit():
        group = 2  # 2. กลุ่ม "ตัวเลขธรรมดา"
        num_part = person_id
    else:
        group = 3  # 3. กลุ่มอื่นๆ
        num_part = person_id
    
    # พยายามแปลงส่วนที่เหลือเป็นตัวเลข
    try:
        sort_val = int(num_part)
        value_type = 0 # 0. เป็นตัวเลข (มาก่อน)
    except ValueError:
        sort_val = num_part # เป็นข้อความ
        value_type = 1 # 1. เป็นข้อความ (มาทีหลัง)

    return (group, value_type, sort_val)

def get_filter_name(verified_status):
    """
    แปลงค่า status เป็นข้อความสำหรับตั้งชื่อไฟล์
    """
    if verified_status == '0':
        return "ยังไม่รายงานตัว"
    elif verified_status == '1':
        return "รายงานตัวแล้ว"
    elif verified_status == '2':
        return "อยู่ในห้องพิธี"
    else: # None
        return "ทั้งหมด"

class ResetDatabase(APIView):
    def post(self, request):
        try:
            with transaction.atomic():
                # 1. บันทึกจำนวนข้อมูลก่อนลบ (Option)
                total_records = Person.objects.count()
                
                # 2. ลบข้อมูลทั้งหมด
                Person.objects.all().delete()
                
                # 3. รีเซ็ต AUTO_INCREMENT (MySQL/MariaDB)
                reset_auto_increment = False
                if 'mysql' in connection.settings_dict['ENGINE']:
                    cursor = connection.cursor()
                    table_name = Person._meta.db_table
                    cursor.execute(f"ALTER TABLE {table_name} AUTO_INCREMENT = 1;")
                    reset_auto_increment = True
                
                # 4. บันทึก Log
                log_details = (
                    f"รีเซ็ตฐานข้อมูล | ลบข้อมูลทั้งหมด {total_records} รายการ"
                )
                
                Log.objects.create(
                    action='Reset',
                    model='Database',
                    details=log_details,
                    record_id=None,
                    user=request.user,
                    user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
                )
                broadcast_ws("reset")
                broadcast_stats_update()
                return Response(
                    {'success': 'รีเซ็ตฐานข้อมูลสำเร็จ'}, 
                    status=status.HTTP_200_OK
                )
        except Exception as e:
            # บันทึก Log กรณี error
            Log.objects.create(
                action='Reset',
                model='Database',
                details=f"รีเซ็ตล้มเหลว: {str(e)}",
                record_id=None,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            return Response(
                {'error': str(e)}, 
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
        
class ResetLog(APIView):
    def post(self, request):
        try:
            with transaction.atomic():
                # 1. บันทึกจำนวนข้อมูลก่อนลบ (Option)
                total_records = Log.objects.count()
                
                # 2. ลบข้อมูลทั้งหมด
                Log.objects.all().delete()
                
                # 3. รีเซ็ต AUTO_INCREMENT (MySQL/MariaDB)
                reset_auto_increment = False
                if 'mysql' in connection.settings_dict['ENGINE']:
                    cursor = connection.cursor()
                    table_name = Log._meta.db_table
                    cursor.execute(f"ALTER TABLE {table_name} AUTO_INCREMENT = 1;")
                    reset_auto_increment = True
                
                # 4. บันทึก Log
                log_details = (
                    f"รีเซ็ตประวัติ | ลบข้อมูลทั้งหมด {total_records} รายการ"
                )
                
                Log.objects.create(
                    action='Reset',
                    model='Database',
                    details=log_details,
                    record_id=None,
                    user=request.user,
                    user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
                )   
                
                return Response(
                    {'success': 'รีเซ็ตประวัติสำเร็จ'}, 
                    status=status.HTTP_200_OK
                )
        except Exception as e:
            # บันทึก Log กรณี error
            Log.objects.create(
                action='Reset',
                model='Database',
                details=f"รีเซ็ตล้มเหลว: {str(e)}",
                record_id=None,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            return Response(
                {'error': str(e)}, 
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

def file_iterator(file, chunk_size=8192):
    while True:
        chunk = file.read(chunk_size)
        if not chunk:
            break
        yield chunk

class ExportPDFResult(View):
    def get(self, request):
        try:
            BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            FONT_PATH = os.path.join(BASE_DIR, 'fonts', 'THSarabunNew.ttf')
            pdfmetrics.registerFont(TTFont('THSarabun', FONT_PATH))

            buffer = io.BytesIO()
            p = canvas.Canvas(buffer, pagesize=A4)
            width, height = A4

            # ---------------- ฟังก์ชันวาดเส้นตาราง ----------------
            def draw_table_grid(canvas, x_list, y_list):
                canvas.setStrokeColorRGB(0, 0, 0)
                canvas.setLineWidth(1)
                for x in x_list:  # เส้นตั้ง
                    canvas.line(x, min(y_list) - 20, x, max(y_list))
                for y in y_list:  # เส้นนอน
                    canvas.line(min(x_list), y, max(x_list), y)
                canvas.line(min(x_list), min(y_list) - 20, max(x_list), min(y_list) - 20)

            # ---------------- Helper ----------------
            def degree_group(name):
                if 'ดุษฎีบัณฑิต' in name:
                    return 'ป.เอก'
                elif 'มหาบัณฑิต' in name:
                    return 'ป.โท'
                return 'ป.ตรี'

            def is_verified(person):
                return any(getattr(person, f'verified{i}') in [1, 2] for i in range(1, 4))

            persons = Person.objects.all()
            degree_summary = {'ป.ตรี': {'total': 0, 'present': 0},
                              'ป.โท': {'total': 0, 'present': 0},
                              'ป.เอก': {'total': 0, 'present': 0}}
            branch_summary = {}

            for person in persons:
                dg = degree_group(person.degree)
                degree_summary[dg]['total'] += 1
                if is_verified(person):
                    degree_summary[dg]['present'] += 1

                branch = person.degree if person.degree else 'ไม่ระบุ'
                if branch not in branch_summary:
                    branch_summary[branch] = {'total': 0, 'present': 0}
                branch_summary[branch]['total'] += 1
                if is_verified(person):
                    branch_summary[branch]['present'] += 1

            # ---------------- HEADER ----------------
            p.setFont('THSarabun', 25)
            date_str = datetime.now().strftime("%d/%m/%Y")
            p.drawRightString(width - 40, height - 40, f"วันที่ {date_str}")
            p.setFont('THSarabun', 15)
            time_str = datetime.now().strftime("%H:%M")
            p.drawRightString(width - 40, height - 60, f"เวลา {time_str}")

            # ---------------- 1. ใบสรุปผล ----------------
            p.setFont('THSarabun', 25)
            p.drawCentredString(width / 2, height - 80, "ใบสรุปผล")

            col_positions = [60, 200, 320, 440, 530]
            col_widths = [140, 120, 120, 90]
            header_positions = [col_positions[i] + (col_widths[i] / 2) for i in range(len(col_widths))]

            headers = ["ระดับ", "ทั้งหมด", "มา", "ขาด"]
            p.setFont('THSarabun', 20)

            # --- ควบคุมความสูงของแต่ละแถว ---
            row_height = 30
            rows = ['header', 'ป.ตรี', 'ป.โท', 'ป.เอก', 'รวมทั้งหมด']

            # คำนวณตำแหน่ง y ของเส้นตาราง (บน → ล่าง)
            y_start = height - 110
            rows_y = [y_start - (i * row_height) for i in range(len(rows) + 1)]

            # --- วาดหัวตาราง ---
            for i, h in enumerate(headers):
                p.drawCentredString(header_positions[i], rows_y[0] - row_height/2 - 5, h)

            total_all = present_all = 0

            # --- วาดแถวข้อมูล ---
            for idx, degree in enumerate(['ป.ตรี', 'ป.โท', 'ป.เอก']):
                total = degree_summary[degree]['total']
                present = degree_summary[degree]['present']
                absent = total - present
                vertical_center = rows_y[idx+1] - row_height/2 - 5

                p.drawCentredString(header_positions[0], vertical_center, degree)
                p.drawCentredString(header_positions[1], vertical_center, str(total))
                p.drawCentredString(header_positions[2], vertical_center, str(present))
                p.drawCentredString(header_positions[3], vertical_center, str(absent))

                total_all += total
                present_all += present

            # --- รวมทั้งหมด ---
            absent_all = total_all - present_all
            vertical_center = rows_y[4] - row_height/2 - 5
            p.drawCentredString(header_positions[0], vertical_center, "รวมทั้งหมด")
            p.drawCentredString(header_positions[1], vertical_center, str(total_all))
            p.drawCentredString(header_positions[2], vertical_center, str(present_all))
            p.drawCentredString(header_positions[3], vertical_center, str(absent_all))

            # --- วาดเส้นตาราง ---
            draw_table_grid(p, col_positions, rows_y)


            # ---------------- 2. ตารางแต่ละสาขา ----------------
            p.showPage()
            p.setFont('THSarabun', 25)
            p.drawCentredString(width / 2, height - 80, "ตารางพระราชทานปริญญาบัตร 2568")

            col_positions = [40, 240, 320, 400, 480, 560]
            col_widths = [200, 80, 80, 80, 80]
            header_positions = [col_positions[i] + (col_widths[i] / 2) for i in range(len(col_widths))]
            headers = ["ชื่อหลักสูตร", "ทั้งหมด", "มา", "ขาด", "ร้อยละ"]

            def draw_branch_header():
                p.setFont('THSarabun', 14)
                for i, h in enumerate(headers):
                    p.drawCentredString(header_positions[i], height - 130, h)

            draw_branch_header()
            y_position = height - 140
            rows_y = [height - 110]

            for branch, vals in sorted(branch_summary.items()):
                total = vals['total']
                present = vals['present']
                absent = total - present
                percent = (present / total * 100) if total > 0 else 0

                vertical_center = y_position - 15
                p.drawCentredString(header_positions[0], vertical_center, branch)
                p.drawCentredString(header_positions[1], vertical_center, str(total))
                p.drawCentredString(header_positions[2], vertical_center, str(present))
                p.drawCentredString(header_positions[3], vertical_center, str(absent))
                p.drawCentredString(header_positions[4], vertical_center, f"{int(percent)}%")

                rows_y.append(y_position)
                y_position -= 20

                if y_position < 50:
                    draw_table_grid(p, col_positions, rows_y)
                    p.showPage()
                    draw_branch_header()
                    rows_y = [height - 110]
                    y_position = height - 140

            if rows_y:
                draw_table_grid(p, col_positions, rows_y)

            # ---------------- 3. รายชื่อที่ยังไม่รายงานตัว (*** แก้ไขส่วนนี้ ***) ----------------
            p.showPage()
            p.setFont('THSarabun', 25)
            p.drawCentredString(width / 2, height - 80, "รายชื่อที่ยังไม่รายงานตัว")

            col_positions = [100, 175, 325, 500]
            col_widths = [75, 150, 175]
            header_positions = [col_positions[i] + (col_widths[i] / 2) for i in range(len(col_widths))]
            headers = ["เลขที่บัณฑิต", "ชื่อ - สกุล", "ชื่อหลักสูตร"]

            def draw_missing_header():
                p.setFont('THSarabun', 14)
                for i, h in enumerate(headers):
                    p.drawCentredString(header_positions[i], height - 130, h)

            draw_missing_header()
            y_position = height - 140
            rows_y = [height - 110]

            # --- Helper function สำหรับการเรียงลำดับแบบพิเศษ ---
            def get_custom_sort_key(person):
                person_id = str(person.id) # แปลงเป็น string ก่อนเสมอ

                if person_id.startswith('ปท'):
                    group = 1  # 1. กลุ่ม "ปท"
                    num_part = person_id[2:] # เอาตัวเลข/ข้อความ หลัง "ปท"
                elif person_id.startswith('ป'):
                    group = 0  # 0. กลุ่ม "ป" (มาก่อน)
                    num_part = person_id[1:] # เอาตัวเลข/ข้อความ หลัง "ป"
                elif person_id.isdigit():
                    group = 2  # 2. กลุ่ม "ตัวเลขธรรมดา"
                    num_part = person_id
                else:
                    group = 3  # 3. กลุ่มอื่นๆ (ถ้ามี)
                    num_part = person_id
                
                # พยายามแปลงส่วนที่เหลือเป็นตัวเลข
                try:
                    sort_val = int(num_part)
                    value_type = 0 # 0. เป็นตัวเลข (มาก่อน)
                except ValueError:
                    sort_val = num_part # เป็นข้อความ
                    value_type = 1 # 1. เป็นข้อความ (มาทีหลัง)

                # คืนค่าเป็น tuple (Group, Type, Value)
                return (group, value_type, sort_val)
            # --- จบ Helper function ---

            # 1. ดึงข้อมูลทั้งหมดที่ยังไม่ verified
            missing_persons_query = Person.objects.filter(
                verified1=0, verified2=0, verified3=0
            )

            # 2. สั่งเรียงใน Python โดยใช้ฟังก์ชันที่เราสร้างขึ้น
            missing_persons = sorted(list(missing_persons_query), key=get_custom_sort_key)
            
            # --- (จบส่วนแก้ไข) ---

            for person in missing_persons:
                vertical_center = y_position - 15
                p.drawCentredString(header_positions[0], vertical_center, str(person.id))
                p.drawCentredString(header_positions[1], vertical_center, person.name)
                p.drawCentredString(header_positions[2], vertical_center, person.degree or "-")

                rows_y.append(y_position)
                y_position -= 20

                if y_position < 50:
                    draw_table_grid(p, col_positions, rows_y)
                    p.showPage()
                    draw_missing_header()
                    rows_y = [height - 110]
                    y_position = height - 140

            if rows_y:
                draw_table_grid(p, col_positions, rows_y)

            # ---------------- SAVE ----------------
            p.save()
            buffer.seek(0)

            date_str = datetime.now().strftime('%Y%m%d')
            filename = f"รายชื่อสรุป_{date_str}.pdf"
            quoted_filename = quote(filename)

            response = StreamingHttpResponse(file_iterator(buffer), content_type='application/pdf')
            response['Content-Disposition'] = (
                f'attachment; filename="{quoted_filename}"; '
                f"filename*=UTF-8''{quoted_filename}"
            )
            response['Cache-Control'] = 'no-store'
            response['Pragma'] = 'no-cache'
            response['Expires'] = '0'

            Log.objects.create(
                action='Export',
                model='Person',
                details="โหลดไฟล์สรุป PDF",
                record_id=None,
                user=request.user,
                user_nickname=getattr(getattr(request.user, "profile", None), "nickname", "")
            )
            return response

        except Exception as e:
            import traceback
            print("PDF Export Error:", str(e))
            print(traceback.format_exc())
            return HttpResponse(f'เกิดข้อผิดพลาด: {str(e)}', status=500)

class ExportPDF(View):
    def get(self, request):
        try:
            # 1. ดึงข้อมูลและสถานะ
            verified_status = request.GET.get('verified_status', None)
            persons_queryset = get_filtered_persons(request)
            
            # 2. เรียงลำดับใน Python
            persons_list = sorted(list(persons_queryset), key=get_custom_sort_key)

            # ตั้งค่า Font
            BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            FONT_PATH = os.path.join(BASE_DIR, 'fonts', 'THSarabunNew.ttf')
            pdfmetrics.registerFont(TTFont('THSarabun', FONT_PATH))
            buffer = io.BytesIO()
            p = canvas.Canvas(buffer, pagesize=A4)
            width, height = A4
            
            def draw_table_grid(canvas, x_list, y_list):
                canvas.setStrokeColorRGB(0, 0, 0)
                canvas.setLineWidth(1)
                for x in x_list:
                    canvas.line(x, min(y_list) - 20, x, max(y_list))
                for y in y_list:
                    canvas.line(min(x_list), y, max(x_list), y)
                canvas.line(min(x_list), min(y_list) - 20, max(x_list), min(y_list) - 20)
            
            # 3. ตั้งค่าตำแหน่งคอลัมน์ (สำหรับ 5 คอลัมน์)
            col_positions = [50, 150, 300, 450, 550, 650] # 5 คอลัมน์ = 6 เส้น
            col_widths = [100, 150, 150, 100, 100]
            
            header_positions = []
            for i in range(len(col_widths)):
                center_x = col_positions[i] + (col_widths[i] / 2)
                header_positions.append(center_x)
            
            # ข้อมูลส่วนหัว (วันที่/เวลา)
            p.setFont('THSarabun', 25)
            date_str = datetime.now().strftime("%d/%m/%Y")
            p.drawRightString(width - 40, height - 40, f"วันที่ {date_str}")
            p.setFont('THSarabun', 15)
            time_str = datetime.now().strftime("%H:%M")
            p.drawRightString(width - 40, height - 60, f"เวลา {time_str}")

            p.setFont('THSarabun', 20)
            p.drawCentredString(width / 2, 780, "รายชื่อบัณฑิต")

            # 4. เขียนหัวตาราง (สำหรับ 5 คอลัมน์)
            p.setFont('THSarabun', 14)
            headers = ["เลขที่บัณฑิต", "ชื่อ - สกุล", "ชื่อหลักสูตร", "รหัส RFID", "สถานะรายงานตัว"]
            for i, header in enumerate(headers):
                p.drawCentredString(header_positions[i], 735, header)

            rows_y = [750]
            y_position = 730
            
            # 5. วาดข้อมูล (สำหรับ 5 คอลัมน์)
            for person in persons_list:
                vertical_center = y_position - 15
                
                p.drawCentredString(header_positions[0], vertical_center, str(person.id))
                p.drawCentredString(header_positions[1], vertical_center, person.name)
                p.drawCentredString(header_positions[2], vertical_center, person.degree or "-")
                p.drawCentredString(header_positions[3], vertical_center, person.rfid or "-")
                p.drawCentredString(header_positions[4], vertical_center, str(person.verified1))
                
                rows_y.append(y_position)
                y_position -= 20

                if y_position < 50:
                    draw_table_grid(p, col_positions, rows_y)
                    p.showPage()
                    
                    y_position = 800
                    rows_y = []
                    
                    p.setFont('THSarabun', 20)
                    p.drawCentredString(width / 2, 780, "รายชื่อบัณฑิต (ต่อ)")
                    p.setFont('THSarabun', 14)
                    for i, header in enumerate(headers):
                        p.drawCentredString(header_positions[i], 735, header)
                    
                    rows_y = [750]
                    y_position = 730

            if len(rows_y) > 0:
                draw_table_grid(p, col_positions, rows_y)

            p.save()
            buffer.seek(0)

            # 6. ตั้งชื่อไฟล์แบบ Dynamic
            filter_name = get_filter_name(verified_status)
            date_str = datetime.now().strftime('%Y%m%d')
            filename = f"รายชื่อ{filter_name}_{date_str}.pdf"
            quoted_filename = quote(filename)

            response = StreamingHttpResponse(
                file_iterator(buffer), 
                content_type='application/pdf'
            )
            response['Cache-Control'] = 'no-store'
            response['Pragma'] = 'no-cache'
            response['Expires'] = '0'
            response['Content-Disposition'] = (
                f'attachment; filename="{quoted_filename}"; '
                f'filename*=UTF-8\'\'{quoted_filename}'
            )
            response["Access-Control-Expose-Headers"] = "Content-Disposition"
            
            Log.objects.create(
                action='Export',
                model='Person',
                details="โหลดไฟล์เป็น PDF",
                record_id=None,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            return response

        except Exception as e:
            print('PDF Export Error:', str(e))
            import traceback
            traceback.print_exc()
            return JsonResponse({'error': str(e)}, status=500)

class ExportData(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request, format_type):
        try:
            # Debug logging
            logger.info(f"Export request: format_type={format_type}, verified_status={request.GET.get('verified_status', None)}")
            
            # 1. ดึงข้อมูลและสถานะ
            verified_status = request.GET.get('verified_status', None)
            persons_queryset = get_filtered_persons(request)

            # 2. เรียงลำดับใน Python
            persons_list = sorted(list(persons_queryset), key=get_custom_sort_key)

            # 3. สร้าง Dataset ด้วยตนเอง (นี่คือการแก้ Error 400)
            dataset = Dataset()
            dataset.headers = ['เลขที่บัณฑิต', 'ชื่อ - สกุล', 'ชื่อหลักสูตร', 'รหัส RFID', 'สถานะรายงานตัว']
            
            for person in persons_list:
                dataset.append([
                    str(person.id),  # แปลงเป็น string เพื่อป้องกันปัญหา
                    person.name or "-",
                    person.degree or "-", # ใช้ or "-" เผื่อค่าว่าง
                    person.rfid or "-",
                    person.verified1  # แสดงเป็นตัวเลข 0, 1, 2
                ])

            response = None
            format_type = format_type.lower()
            
            # Validate format_type
            if format_type not in ['excel', 'csv']:
                logger.error(f"Invalid format_type: {format_type}")
                return Response(
                    {'error': f'รูปแบบไฟล์ไม่ถูกต้อง: {format_type}'}, 
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # 4. ตั้งชื่อไฟล์แบบ Dynamic
            filter_name = get_filter_name(verified_status)
            date_str = datetime.now().strftime('%Y%m%d')
            base_filename = f"รายชื่อ{filter_name}_{date_str}"

            if format_type == 'excel':
                response = HttpResponse(
                    dataset.xlsx,
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                )
                filename = f'{base_filename}.xlsx'
                response['Content-Disposition'] = f'attachment; filename="{filename}"'
                response['Access-Control-Expose-Headers'] = 'Content-Disposition'
                response['X-Filename'] = filename  # ส่งชื่อไฟล์ผ่าน custom header

            elif format_type == 'csv':
                response = HttpResponse(dataset.csv, content_type='text/csv; charset=utf-8-sig')
                filename = f'{base_filename}.csv'
                response['Content-Disposition'] = f'attachment; filename="{filename}"'
                response['Access-Control-Expose-Headers'] = 'Content-Disposition'
                response['X-Filename'] = filename  # ส่งชื่อไฟล์ผ่าน custom header

            else:
                return Response(
                    {'error': 'รูปแบบไฟล์ไม่ถูกต้อง'}, 
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            Log.objects.create(
                action='Export',
                model='Person',
                details=f"โหลดไฟล์เป็น {format_type}",
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            return response

        except Exception as e:
            import traceback
            traceback.print_exc() # พิมพ์ error ออกมาดูใน console
            logger.error(f"Export error: {str(e)}", exc_info=True)
            return Response(
                {'error': f'Export failed: {str(e)}'}, 
                status=500
            )

class ImportData(APIView):
    def post(self, request):
        file = request.FILES['file']

        # ตรวจสอบขนาดไฟล์ (ไม่เกิน 15MB)
        max_size = 15 * 1024 * 1024  # 15 MB
        if file.size > max_size:
            return Response(
                {'error': 'ไฟล์มีขนาดเกิน 15MB'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        dataset = Dataset()
        resource = PersonResource()

        try:
            # อ่านไฟล์
            if file.name.endswith('.xlsx'):
                dataset.load(file.read(), format='xlsx')
            elif file.name.endswith('.csv'):
                dataset.load(file.read(), format='csv', encoding='utf-8-sig')

            # ตรวจสอบข้อมูล
            if len(dataset) == 0:
                raise ValueError("ไฟล์ที่อัปโหลดว่างเปล่า")
            

            # นำเข้าข้อมูล (เพิ่ม update=True เพื่อให้ทับข้อมูลเก่า)
            result = resource.import_data(
                dataset, 
                dry_run=False, 
                update=True,
                skip_unchanged=False,  # เปลี่ยนเป็น False เพื่อให้อัปเดตทุกครั้ง
                use_bulk=False
            )
            
            
            
            # Log ไปยัง logger ด้วย
            logger.info(f"Import result: {result.totals}")
            logger.info(f"New records: {result.totals.get('new', 0)}")
            logger.info(f"Updated records: {result.totals.get('update', 0)}")
            logger.info(f"Errors: {result.totals.get('error', 0)}")
            logger.info(f"Delete: {result.totals.get('delete', 0)}")
            logger.info(f"Skip: {result.totals.get('skip', 0)}")
            logger.info(f"Invalid: {result.totals.get('invalid', 0)}")
            
            if result.has_errors():
                logger.error("Import errors found:")
                for error in result.row_errors():
                    logger.error(f"Row {error[0]}: {error[1]}")
            
            # แก้ไขการนับจำนวนรายการ
            imported_count = (
                result.totals.get('new', 0)    # ข้อมูลใหม่
                + result.totals.get('update', 0)  # ข้อมูลที่อัปเดต
            )

            # คำนวณสถิติเพิ่มเติม
            all_persons = Person.objects.all()
            total_count = all_persons.count()
            rfid_count = all_persons.exclude(Q(rfid__isnull=True) | Q(rfid='')).count()
            
            # นับจำนวนตามระดับปริญญา
            bachelor_count = 0
            master_count = 0
            doctor_count = 0
            
            for person in all_persons:
                if person.degree:
                    if 'ดุษฎีบัณฑิต' in person.degree:
                        doctor_count += 1
                    elif 'มหาบัณฑิต' in person.degree:
                        master_count += 1
                    else:
                        bachelor_count += 1

            # บันทึก Log
            Log.objects.create(
                action='Import',
                model='Person',
                details=f"นำเข้าฐานข้อมูล {imported_count} รายการ ( ใหม่ {result.totals.get('new', 0)} อัปเดต {result.totals.get('update', 0)} )",
                record_id=None,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            broadcast_ws("upload")
            broadcast_stats_update()
            
            return Response(
                {
                    'success': f'นำเข้าข้อมูลสำเร็จ {imported_count} รายการ',
                    'stats': {
                        'imported_count': imported_count,
                        'new_count': result.totals.get('new', 0),
                        'updated_count': result.totals.get('update', 0),
                        'total_count': total_count,
                        'rfid_count': rfid_count,
                        'bachelor_count': bachelor_count,
                        'master_count': master_count,
                        'doctor_count': doctor_count
                    }
                }, 
                status=status.HTTP_201_CREATED
            )
            
        except Exception as e:
            
            Log.objects.create(
                action='Import',
                model='Person',
                details=f"นำเข้าข้อมูลล้มเหลว: {str(e)}",
                record_id=None,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            logger.error(f"Import failed: {str(e)}", exc_info=True)
            return Response(
                {'error': str(e)}, 
                status=status.HTTP_400_BAD_REQUEST
            )

class StatsView(APIView):
    def get(self, request):
        from collections import Counter
        verified_counter = Counter()
        persons = Person.objects.all()
        total = persons.count()

        for person in persons:
            verified_with_time = []
            for i in range(1, 4):
                value = getattr(person, f'verified{i}', 0)
                timestamp = getattr(person, f'verified_updated_at{i}', None)

                # ถ้า timestamp ไม่มี ให้ใช้วันที่เก่ามากๆ แทน เพื่อให้ไม่ถูกเลือกก่อน timestamp อื่น
                if not timestamp:
                    timestamp = datetime.min.replace(tzinfo=timezone.utc)
                verified_with_time.append((timestamp, value))

            if verified_with_time:
                # หาค่าล่าสุดจาก timestamp ที่ใหม่ที่สุด
                latest_value = sorted(verified_with_time, key=lambda x: x[0], reverse=True)[0][1]
                verified_counter[latest_value] += 1
    
        stats = {
            'total': total,
            'not_checked_in': verified_counter[0],      # ยังไม่รายงานตัว (สถานะ 0)
            'in_checkin_room': verified_counter[1],     # รายงานตัวแล้ว (สถานะ 1)
            'in_graduation_room': verified_counter[2],  # อยู่ในห้องพิธี (สถานะ 2)
        }
        return Response(stats, status=200)

class PersonList(APIView):
    def get(self, request):
        persons = Person.objects.all()
        serializer = PersonSerializer(persons, many=True)
        # ตรวจสอบข้อมูลก่อนส่ง response
        safe_data = []
        for item in serializer.data:
            safe_item = {
                k: v for k, v in item.items() 
                if isinstance(v, (str, int, float, bool, type(None)))
            }
            safe_data.append(safe_item)
        return Response(safe_data, status=status.HTTP_200_OK)
    
    def post(self, request):
        serializer = PersonSerializer(data=request.data)
        if serializer.is_valid():
            instance = serializer.save()  # เก็บ instance ที่สร้าง
            Log.objects.create(
                action='Add',
                model='Person',
                details=f"เพิ่มข้อมูล: {instance.name}",
                record_id=instance.id,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            if settings.USE_CHANNEL:
                broadcast_to_crud01({
                    'action': 'add',
                    'id': instance.id,
                    'fields': {
                        'name': instance.name,
                        'nisit': instance.nisit,
                        'degree': instance.degree,
                        'seat': instance.seat,
                        'verified1': instance.verified1,
                        'verified2': instance.verified2,
                        'verified3': instance.verified3,
                        'rfid': instance.rfid,
                    }
                })
                broadcast_stats_update()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def put(self, request):
        ids = request.data.get('ids', [])
        verified = request.data.get('verified')
        verified_field = request.data.get('verified_field')

        if not ids or not verified_field:
            return Response({'error': 'ข้อมูลไม่ครบ'}, status=400)

        now = timezone.now()
        new_val = int(verified)

        updated_field = verified_field.replace('verified', 'verified_updated_at')

        # --- อัปเดตครั้งเดียว ---
        updated_count = Person.objects.filter(id__in=ids).exclude(**{verified_field: new_val}).update(
            **{
                verified_field: new_val,
                updated_field: now
            }
        )

        # --- ส่ง broadcast ทีเดียว ---
        if updated_count > 0 and settings.USE_CHANNEL:
            broadcast_to_crud01({
                'action': 'bulk_update',
                'ids': ids,
                'fields': {
                    verified_field: new_val,
                    updated_field: now.isoformat(),
                }
            })

            Log.objects.create(
                action='Edit',
                model='Person',
                details=f"[ID: {','.join(map(str, ids))}] อัปเดตเป็น {new_val}",
                user=request.user,
                user_nickname=getattr(getattr(request.user, 'profile', None), 'nickname', '')
            )

            broadcast_stats_update()

        return Response({
            'updated_count': updated_count,
            'updated_ids': ids,
        }, status=200)

    def delete(self, request):
        ids = request.data.get('ids', [])
        
        if not ids:
            return Response({'error': 'No IDs provided'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            with transaction.atomic():
                persons = Person.objects.filter(id__in=ids)
                # เก็บ ID ก่อนลบ
                ids_str = ','.join(str(p.id) for p in persons)
                Log.objects.create(
                    action='Delete',
                    model='Person',
                    details=f"[ID: {ids_str}] ลบข้อมูลแบบกลุ่ม",
                    record_id=None,
                    user=request.user,
                    user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
                )
                persons.delete()
                
                # ส่ง WebSocket สำหรับแต่ละ ID
                if settings.USE_CHANNEL:
                    for id in ids:
                        broadcast_to_crud01({
                            'action': 'delete',
                            'id': id,
                        })
                    broadcast_stats_update()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

class PersonDetail(APIView):
    def get(self, request, pk):
        try:
            person = Person.objects.get(pk=pk)
        except Person.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)

        serializer = PersonSerializer(person)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def delete(self, request, pk):
        try:
            person = Person.objects.get(pk=pk)
            person_id = person.id  # เก็บ ID ก่อนลบ
            
            Log.objects.create(
                action='Delete',
                model='Person',
                details=f"ลบข้อมูลของ {person.name}",
                record_id=person.id,
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            
            person.delete()
            
            # ส่ง WebSocket action delete
            if settings.USE_CHANNEL:
                broadcast_to_crud01({
                    'action': 'delete',
                    'id': person_id,
                })
                broadcast_stats_update()
            return Response(status=status.HTTP_204_NO_CONTENT)
        except Person.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)
    
    def put(self, request, pk):
        try:
            person = Person.objects.get(pk=pk)
            original_data = person_to_dict(person)
            serializer = PersonSerializer(person, data=request.data)
            if serializer.is_valid():
                serializer.save()
                person.refresh_from_db()
                changes = []
                for field in ['name', 'degree', 'seat', 'verified1', 'verified2', 'verified3', 'rfid']:
                    old_val = original_data[field]
                    new_val = getattr(person, field)
                    if old_val != new_val:
                        changes.append(f"{field}::{old_val}::{new_val}")
                if changes:
                    log_message = " | ".join(changes)
                    Log.objects.create(
                        action='Edit',
                        model='Person',
                        details=log_message, 
                        record_id=person.id,
                        user=request.user,
                        user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
                    )
                if settings.USE_CHANNEL:
                    fields = person_to_dict(person)
                    fields = convert_datetime_fields(fields, [
                        'verified_updated_at1', 'verified_updated_at2', 'verified_updated_at3'
                    ])
                    broadcast_to_crud01({
                        'action': 'update',
                        'id': person.id,
                        'fields': fields,
                    })
                    broadcast_stats_update()
                return Response(serializer.data, status=status.HTTP_200_OK)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        except Person.DoesNotExist:
            return Response(status=status.HTTP_404_NOT_FOUND)
        
class RFIDReceive1(APIView):
    parser_classes = [JSONParser]
    def post(self, request):
        try:
            simulated_tags = request.data.get('tags', [])
            scanner_type = request.data.get('scanner_type')  # 'in' or 'out'
            scanner_id = 1  # ตายตัวเป็น 1

            if not simulated_tags or scanner_type not in ['in', 'out']:
                return Response(
                    {'error': 'Missing tags or invalid scanner_type (in/out)'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            results = []

            for tag in simulated_tags:
                epc = tag.get('epc')
                if not epc:
                    continue
                try:
                    # ✅ Step 1: ดึงข้อมูล person เพื่อตรวจสอบสถานะปัจจุบันก่อน
                    person = Person.objects.get(rfid=epc)

                    verified_field = f"verified{scanner_id}"
                    time_field = f"verified_updated_at{scanner_id}"
                    current_status = getattr(person, verified_field, 0)
                    verified_value = 2 if scanner_type == 'out' else 1

                    if current_status == verified_value:
                        results.append(f"rfid: {epc} name: {person.name} status: แท็กนี้ถูกแสกนแล้ว")
                    else:
                        now = timezone.now()
                        
                        # ✅ Step 2: ใช้ .update() เพื่อสั่งให้ฐานข้อมูลอัปเดตโดยตรง
                        #    วิธีนี้แน่นอนและมีประสิทธิภาพกว่าการใช้ .save()
                        Person.objects.filter(pk=person.pk).update(**{
                            verified_field: verified_value,
                            time_field: now
                        })

                        if settings.USE_CHANNEL:
                            # ✅ Step 3: อัปเดต object ในหน่วยความจำตาม เพื่อส่งข้อมูลที่ถูกต้องผ่าน WebSocket
                            setattr(person, verified_field, verified_value)
                            setattr(person, time_field, now)
                            
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person.id,
                                'fields': person_to_dict(person),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()

                        results.append(f"rfid: {epc} name: {person.name} status: อัปเดตสถานะสำเร็จ")

                except Person.DoesNotExist:
                    # ส่วนนี้ทำงานถูกต้องอยู่แล้ว ไม่ต้องแก้ไข
                    person_with_empty_rfid = Person.objects.filter(Q(rfid__isnull=True) | Q(rfid='')).first()

                    if not person_with_empty_rfid:
                        results.append(f"rfid: {epc} name: null status: ไม่พบข้อมูลในระบบ")
                    else:
                        person_with_empty_rfid.rfid = epc
                        verified_field = f"verified{scanner_id}"
                        time_field = f"verified_updated_at{scanner_id}"
                        verified_value = 2 if scanner_type == 'out' else 1
                        setattr(person_with_empty_rfid, verified_field, verified_value)
                        setattr(person_with_empty_rfid, time_field, timezone.now())
                        person_with_empty_rfid.save()

                        if settings.USE_CHANNEL:
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person_with_empty_rfid.id,
                                'fields': person_to_dict(person_with_empty_rfid),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()
                        results.append(f"epc: {epc} name: {person_with_empty_rfid.name} status: เพิ่มรหัส RFID สำเร็จและอัปเดตสถานะแล้ว")

            return Response({'results': results}, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

class RFIDReceive2(APIView):
    parser_classes = [JSONParser]
    def post(self, request):
        try:
            simulated_tags = request.data.get('tags', [])
            scanner_type = request.data.get('scanner_type')  # 'in' or 'out'
            scanner_id = 2  # ตายตัวเป็น 2

            if not simulated_tags or scanner_type not in ['in', 'out']:
                return Response(
                    {'error': 'Missing tags or invalid scanner_type (in/out)'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            results = []

            for tag in simulated_tags:
                epc = tag.get('epc')
                if not epc:
                    continue
                try:
                    # ✅ Step 1: ดึงข้อมูล person เพื่อตรวจสอบสถานะปัจจุบันก่อน
                    person = Person.objects.get(rfid=epc)

                    verified_field = f"verified{scanner_id}"
                    time_field = f"verified_updated_at{scanner_id}"
                    current_status = getattr(person, verified_field, 0)
                    verified_value = 2 if scanner_type == 'out' else 1

                    if current_status == verified_value:
                        results.append(f"rfid: {epc} name: {person.name} status: แท็กนี้ถูกแสกนแล้ว")
                    else:
                        now = timezone.now()
                        
                        # ✅ Step 2: ใช้ .update() เพื่อสั่งให้ฐานข้อมูลอัปเดตโดยตรง
                        #    วิธีนี้แน่นอนและมีประสิทธิภาพกว่าการใช้ .save()
                        Person.objects.filter(pk=person.pk).update(**{
                            verified_field: verified_value,
                            time_field: now
                        })

                        if settings.USE_CHANNEL:
                            # ✅ Step 3: อัปเดต object ในหน่วยความจำตาม เพื่อส่งข้อมูลที่ถูกต้องผ่าน WebSocket
                            setattr(person, verified_field, verified_value)
                            setattr(person, time_field, now)
                            
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person.id,
                                'fields': person_to_dict(person),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()

                        results.append(f"rfid: {epc} name: {person.name} status: อัปเดตสถานะสำเร็จ")

                except Person.DoesNotExist:
                    # ส่วนนี้ทำงานถูกต้องอยู่แล้ว ไม่ต้องแก้ไข
                    person_with_empty_rfid = Person.objects.filter(Q(rfid__isnull=True) | Q(rfid='')).first()

                    if not person_with_empty_rfid:
                        results.append(f"rfid: {epc} name: null status: ไม่พบข้อมูลในระบบ")
                    else:
                        person_with_empty_rfid.rfid = epc
                        verified_field = f"verified{scanner_id}"
                        time_field = f"verified_updated_at{scanner_id}"
                        verified_value = 2 if scanner_type == 'out' else 1
                        setattr(person_with_empty_rfid, verified_field, verified_value)
                        setattr(person_with_empty_rfid, time_field, timezone.now())
                        person_with_empty_rfid.save()

                        if settings.USE_CHANNEL:
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person_with_empty_rfid.id,
                                'fields': person_to_dict(person_with_empty_rfid),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()
                        results.append(f"epc: {epc} name: {person_with_empty_rfid.name} status: เพิ่มรหัส RFID สำเร็จและอัปเดตสถานะแล้ว")

            return Response({'results': results}, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

class RFIDReceive3(APIView):
    parser_classes = [JSONParser]
    def post(self, request):
        try:
            simulated_tags = request.data.get('tags', [])
            scanner_type = request.data.get('scanner_type')  # 'in' or 'out'
            scanner_id = 3  # ตายตัวเป็น 3

            if not simulated_tags or scanner_type not in ['in', 'out']:
                return Response(
                    {'error': 'Missing tags or invalid scanner_type (in/out)'},
                    status=status.HTTP_400_BAD_REQUEST
                )

            results = []

            for tag in simulated_tags:
                epc = tag.get('epc')
                if not epc:
                    continue
                try:
                    # ✅ Step 1: ดึงข้อมูล person เพื่อตรวจสอบสถานะปัจจุบันก่อน
                    person = Person.objects.get(rfid=epc)

                    verified_field = f"verified{scanner_id}"
                    time_field = f"verified_updated_at{scanner_id}"
                    current_status = getattr(person, verified_field, 0)
                    verified_value = 2 if scanner_type == 'out' else 1

                    if current_status == verified_value:
                        results.append(f"rfid: {epc} name: {person.name} status: แท็กนี้ถูกแสกนแล้ว")
                    else:
                        now = timezone.now()
                        
                        # ✅ Step 2: ใช้ .update() เพื่อสั่งให้ฐานข้อมูลอัปเดตโดยตรง
                        #    วิธีนี้แน่นอนและมีประสิทธิภาพกว่าการใช้ .save()
                        Person.objects.filter(pk=person.pk).update(**{
                            verified_field: verified_value,
                            time_field: now
                        })

                        if settings.USE_CHANNEL:
                            # ✅ Step 3: อัปเดต object ในหน่วยความจำตาม เพื่อส่งข้อมูลที่ถูกต้องผ่าน WebSocket
                            setattr(person, verified_field, verified_value)
                            setattr(person, time_field, now)
                            
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person.id,
                                'fields': person_to_dict(person),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()

                        results.append(f"rfid: {epc} name: {person.name} status: อัปเดตสถานะสำเร็จ")

                except Person.DoesNotExist:
                    # ส่วนนี้ทำงานถูกต้องอยู่แล้ว ไม่ต้องแก้ไข
                    person_with_empty_rfid = Person.objects.filter(Q(rfid__isnull=True) | Q(rfid='')).first()

                    if not person_with_empty_rfid:
                        results.append(f"rfid: {epc} name: null status: ไม่พบข้อมูลในระบบ")
                    else:
                        person_with_empty_rfid.rfid = epc
                        verified_field = f"verified{scanner_id}"
                        time_field = f"verified_updated_at{scanner_id}"
                        verified_value = 2 if scanner_type == 'out' else 1
                        setattr(person_with_empty_rfid, verified_field, verified_value)
                        setattr(person_with_empty_rfid, time_field, timezone.now())
                        person_with_empty_rfid.save()

                        if settings.USE_CHANNEL:
                            broadcast_to_crud01({
                                'action': 'update',
                                'id': person_with_empty_rfid.id,
                                'fields': person_to_dict(person_with_empty_rfid),
                                'scanner_type': scanner_type,
                            })
                            broadcast_stats_update()
                        results.append(f"epc: {epc} name: {person_with_empty_rfid.name} status: เพิ่มรหัส RFID สำเร็จและอัปเดตสถานะแล้ว")

            return Response({'results': results}, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        
class LogPagination(PageNumberPagination):
    page_size = 100  # เพิ่มจาก 5 เป็น 100 เพื่อแสดงคอมเมนต์มากขึ้น

class LogList(generics.ListAPIView):
    serializer_class = LogSerializer
    pagination_class = LogPagination
    
    def get_queryset(self):
        # กรองข้อมูลที่อาจมี timestamp เป็น null
        return Log.objects.exclude(timestamp__isnull=True).order_by('-timestamp')

class LogCreateView(APIView):
    def post(self, request):
        # เพิ่มข้อมูล user และ user_nickname จาก request
        data = request.data.copy()
        if request.user.is_authenticated:
            data['user'] = request.user.id
            # ใช้ nickname จาก Profile หรือ username เป็น fallback
            try:
                profile = request.user.profile
                data['user_nickname'] = profile.nickname or request.user.username
            except:
                data['user_nickname'] = request.user.username
        
        serializer = LogSerializer(data=data)
        if serializer.is_valid():
            log = serializer.save()

            if log.action == "comment":
                if settings.USE_CHANNEL:
                    broadcast_ws("comment", {
                        "comment": log.details,
                        "time": log.timestamp.isoformat(),
                        "user_nickname": log.user_nickname or "ผู้ใช้ไม่ระบุชื่อ"
                    })

            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
    
def person_to_dict(person):
    return {
        'name': person.name,
        'nisit': person.nisit,
        'degree': person.degree,
        'seat': person.seat,
        'verified1': person.verified1,
        'verified2': person.verified2,
        'verified3': person.verified3,
        'verified_updated_at1': datetime_to_str(person.verified_updated_at1),
        'verified_updated_at2': datetime_to_str(person.verified_updated_at2),
        'verified_updated_at3': datetime_to_str(person.verified_updated_at3),
        'rfid': person.rfid,
    }

def datetime_to_str(dt):
    if dt is None:
        return None
    return dt.isoformat()

def convert_datetime_fields(data: dict, fields: list):
    for f in fields:
        if f in data and isinstance(data[f], datetime):
            data[f] = data[f].isoformat()
    return data

class ExportSeatMap(APIView):
    """
    Export แผนที่นั่งเป็น Excel โดยแสดง:
    - แกน Y = เลขแถว
    - แกน X = เลขที่นั่ง (พร้อมทางเดิน)
    - เซลล์ = ลำดับที่บัณฑิต
    - สีเขียว = มารายงานตัว, สีแดง = ยังไม่มา
    - ทางเดิน = คอลัมน์ว่าง
    """
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            import json
            
            # ดึงข้อมูล layout จาก query parameters
            rows_param = request.GET.get('rows')
            seats_per_row_param = request.GET.get('seatsPerRow')
            seats_per_side_a_param = request.GET.get('seatsPerSideA')
            seats_per_side_b_param = request.GET.get('seatsPerSideB')
            pillars_param = request.GET.get('pillars', '[]')
            
            # ดึงข้อมูลทั้งหมด
            persons = Person.objects.all()
            
            # กำหนดค่าเริ่มต้น
            if rows_param and seats_per_side_a_param:
                SEATS_PER_ROW = int(seats_per_row_param)
                SEATS_PER_SIDE_A = int(seats_per_side_a_param)
                SEATS_PER_SIDE_B = int(seats_per_side_b_param)
                pillars = json.loads(pillars_param)
                max_rows_config = int(rows_param)
            else:
                SEATS_PER_ROW = 70
                SEATS_PER_SIDE_A = 35
                SEATS_PER_SIDE_B = 35
                pillars = []
                max_rows_config = 70
            
            # หาจำนวนแถวที่มีข้อมูลจริง
            max_seat = max([p.seat for p in persons if p.seat], default=0)
            total_rows = min(
                (max_seat // SEATS_PER_ROW) + (1 if max_seat % SEATS_PER_ROW > 0 else 0),
                max_rows_config
            ) if max_seat > 0 else 0
            
            # สร้าง workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "แผนที่นั่ง"
            
            # สร้าง dictionary สำหรับค้นหาบัณฑิตตามที่นั่ง
            seat_map = {}
            for person in persons:
                if person.seat:
                    seat_map[int(person.seat)] = person
            
            # กำหนดสีและ style
            green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            red_fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
            header_fill = PatternFill(start_color="B0E0E6", end_color="B0E0E6", fill_type="solid")
            aisle_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")  # สีเทาสำหรับทางเดิน
            pillar_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")  # สีเทาเข้มสำหรับเสา
            header_font = Font(bold=True, size=12)
            center_alignment = Alignment(horizontal='center', vertical='center')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # ฟังก์ชันสร้าง layout สำหรับแต่ละแถว
            def get_row_layout(row_idx):
                # เริ่มต้นด้วยที่นั่งทั้งหมด
                layout = ['seat'] * SEATS_PER_ROW
                
                # แทรกเสา (pillar)
                for pillar in pillars:
                    if pillar['row'] == row_idx:
                        side = pillar['side']
                        index = pillar['index']
                        length = pillar.get('length', 1)
                        
                        # คำนวณตำแหน่งจริง
                        if side == 'left':
                            start_idx = index
                        else:  # right
                            start_idx = SEATS_PER_SIDE_A + index
                        
                        # แทนที่เป็น pillar
                        for i in range(length):
                            if start_idx + i < len(layout):
                                layout[start_idx + i] = 'pillar'
                
                return layout
            
            # สร้างส่วนหัวคอลัมน์
            ws.cell(row=1, column=1, value="แถว\\ที่นั่ง")
            ws.cell(row=1, column=1).fill = header_fill
            ws.cell(row=1, column=1).font = header_font
            ws.cell(row=1, column=1).alignment = center_alignment
            ws.cell(row=1, column=1).border = thin_border
            
            # นับตำแหน่งคอลัมน์ใน Excel (รวมทางเดิน)
            excel_col = 2
            seat_counter = 1
            
            # สร้างหัวคอลัมน์สำหรับฝั่งซ้าย
            for i in range(SEATS_PER_SIDE_A):
                ws.cell(row=1, column=excel_col, value=seat_counter)
                ws.cell(row=1, column=excel_col).fill = header_fill
                ws.cell(row=1, column=excel_col).font = header_font
                ws.cell(row=1, column=excel_col).alignment = center_alignment
                ws.cell(row=1, column=excel_col).border = thin_border
                ws.column_dimensions[ws.cell(row=1, column=excel_col).column_letter].width = 5
                excel_col += 1
                seat_counter += 1
            
            # ทางเดินกลาง
            ws.cell(row=1, column=excel_col, value="ทางเดิน")
            ws.cell(row=1, column=excel_col).fill = aisle_fill
            ws.cell(row=1, column=excel_col).font = header_font
            ws.cell(row=1, column=excel_col).alignment = center_alignment
            ws.cell(row=1, column=excel_col).border = thin_border
            ws.column_dimensions[ws.cell(row=1, column=excel_col).column_letter].width = 8
            aisle_col = excel_col
            excel_col += 1
            
            # สร้างหัวคอลัมน์สำหรับฝั่งขวา
            for i in range(SEATS_PER_SIDE_B):
                ws.cell(row=1, column=excel_col, value=seat_counter)
                ws.cell(row=1, column=excel_col).fill = header_fill
                ws.cell(row=1, column=excel_col).font = header_font
                ws.cell(row=1, column=excel_col).alignment = center_alignment
                ws.cell(row=1, column=excel_col).border = thin_border
                ws.column_dimensions[ws.cell(row=1, column=excel_col).column_letter].width = 5
                excel_col += 1
                seat_counter += 1
            
            # เติมข้อมูลแต่ละแถว
            for row_num in range(1, total_rows + 1):
                row_layout = get_row_layout(row_num - 1)
                
                # คอลัมน์แรก = เลขแถว
                ws.cell(row=row_num + 1, column=1, value=f"แถว {row_num}")
                ws.cell(row=row_num + 1, column=1).fill = header_fill
                ws.cell(row=row_num + 1, column=1).font = header_font
                ws.cell(row=row_num + 1, column=1).alignment = center_alignment
                ws.cell(row=row_num + 1, column=1).border = thin_border
                
                excel_col = 2
                seat_in_row = 1
                
                # เติมข้อมูลแต่ละที่นั่ง
                for i in range(SEATS_PER_ROW):
                    # ถ้าถึงทางเดินกลาง ให้ข้ามคอลัมน์
                    if i == SEATS_PER_SIDE_A:
                        ws.cell(row=row_num + 1, column=aisle_col, value="")
                        ws.cell(row=row_num + 1, column=aisle_col).fill = aisle_fill
                        ws.cell(row=row_num + 1, column=aisle_col).border = thin_border
                        excel_col = aisle_col + 1
                    
                    cell = ws.cell(row=row_num + 1, column=excel_col if i < SEATS_PER_SIDE_A else excel_col)
                    
                    if row_layout[i] == 'pillar':
                        # เสา
                        cell.value = "▓"
                        cell.fill = pillar_fill
                    else:
                        # ที่นั่ง
                        seat_number = (row_num - 1) * SEATS_PER_ROW + (i + 1)
                        
                        if seat_number in seat_map:
                            person = seat_map[seat_number]
                            cell.value = person.id
                            
                            # กำหนดสีตามสถานะ
                            is_verified = any([
                                person.verified1 in [1, 2],
                                person.verified2 in [1, 2],
                                person.verified3 in [1, 2]
                            ])
                            
                            cell.fill = green_fill if is_verified else red_fill
                        else:
                            cell.value = ""
                    
                    cell.alignment = center_alignment
                    cell.border = thin_border
                    excel_col += 1
            
            # ปรับความกว้างคอลัมน์แรก
            ws.column_dimensions['A'].width = 12
            
            # บันทึกไฟล์
            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            # ตั้งชื่อไฟล์
            date_str = datetime.now().strftime('%Y%m%d')
            filename = f"แผนที่นั่ง_{date_str}.xlsx"
            
            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            response['Access-Control-Expose-Headers'] = 'Content-Disposition'
            
            # Log การ export
            Log.objects.create(
                action='Export',
                model='Person',
                details="โหลดแผนที่นั่งเป็น Excel (พร้อมทางเดิน)",
                user=request.user,
                user_nickname=request.user.profile.nickname if hasattr(request.user, 'profile') else ''
            )
            
            return response
            
        except Exception as e:
            import traceback
            traceback.print_exc()
            return Response(
                {'error': str(e)}, 
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )